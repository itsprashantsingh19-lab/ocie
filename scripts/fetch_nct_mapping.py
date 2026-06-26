import openpyxl, requests, json, time, re, urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

XLSX_PATH = r"C:\Users\ipras\AppData\Local\Temp\opencode\ocie\data\Current Treatment mapping(NCCN_ASCO) for NSCLC.xlsx"
OUTPUT_PATH = r"C:\Users\ipras\AppData\Local\Temp\opencode\ocie\data\nct_mapping.json"

# Drug aliases: xlsx name -> search terms
SEARCH_OVERRIDES = {
    "Ado-trastuzumab emtansine (T-DM1)": "trastuzumab emtansine",
    "Amivantamab-vmjw": "amivantamab",
    "Amivantamab + Hyaluronidase-lpuj": "amivantamab",
    "Amivantamab-vmjw + Carboplatin + Pemetrexed": "amivantamab",
    "Datopotamab deruxtecan-dlnk": "datopotamab deruxtecan",
    "Fam-trastuzumab deruxtecan-nxki": "trastuzumab deruxtecan",
    "Pembrolizumab + Carboplatin + Nab-paclitaxel": "pembrolizumab",
    "Pembrolizumab + Carboplatin + Paclitaxel": "pembrolizumab",
    "Pembrolizumab + Carboplatin + Pemetrexed": "pembrolizumab",
    "Pembrolizumab + Cisplatin + Nab-paclitaxel": "pembrolizumab",
    "Pembrolizumab + Cisplatin + Paclitaxel": "pembrolizumab",
    "Pembrolizumab + Cisplatin + Pemetrexed": "pembrolizumab",
    "Nivolumab + Ipilimumab": "nivolumab ipilimumab",
    "Nivolumab + Ipilimumab + 2 cycles platinum chemotherapy": "nivolumab ipilimumab",
    "Carboplatin + Osimertinib + Pemetrexed": "osimertinib",
    "Cisplatin + Osimertinib + Pemetrexed": "osimertinib",
    "Lazertinib + Amivantamab-vmjw": "lazertinib amivantamab",
    "Dabrafenib + Trametinib": "dabrafenib trametinib",
    "Encorafenib + Binimetinib": "encorafenib binimetinib",
    "Docetaxel +/- Ramucirumab": "docetaxel ramucirumab",
    "Tremelimumab + Durvalumab + Carboplatin + Paclitaxel": "tremelimumab durvalumab",
    "Tremelimumab + Durvalumab + Carboplatin + Pemetrexed": "tremelimumab durvalumab",
    "Pembrolizumab + Carboplatin + Pemetrexed": "pembrolizumab",
    "Pembrolizumab + Cisplatin + Pemetrexed": "pembrolizumab",
    "Cisplatin + Pemetrexed": "cisplatin pemetrexed",
    "Carboplatin + Pemetrexed": "carboplatin pemetrexed",
    "Carboplatin + Paclitaxel": "carboplatin paclitaxel",
    "Carboplatin + Gemcitabine": "carboplatin gemcitabine",
}

def normalize_term(name):
    name = name.strip()
    if name in SEARCH_OVERRIDES:
        return SEARCH_OVERRIDES[name]
    # Remove trailing suffixes in parens like "(T-DM1)"
    name = re.sub(r'\s*\([^)]*\)\s*', ' ', name)
    # Remove suffixes like -vmjw, -dlnk, -nxki, -lpuj
    name = re.sub(r'-\w{4}\b', '', name)
    name = name.strip()
    return name.lower()

def search_nct(term, max_results=5):
    query = f"NSCLC AND {term}"
    url = f"https://clinicaltrials.gov/api/v2/studies?query.term={requests.utils.quote(query)}&pageSize={max_results}&sort=LastUpdatePostDate"
    try:
        r = requests.get(url, verify=False, timeout=15)
        if r.status_code != 200:
            return [], f"HTTP {r.status_code}"
        data = r.json()
        studies = data.get("studies", [])
        results = []
        for s in studies:
            ps = s["protocolSection"]
            nct_id = ps["identificationModule"]["nctId"]
            title = ps["identificationModule"].get("briefTitle", "")
            phases = ps.get("designModule", {}).get("phases", [])
            status = ps.get("statusModule", {}).get("overallStatus", "")
            start_date = ps.get("statusModule", {}).get("startDateStruct", {}).get("date", "")
            completion_date = ps.get("statusModule", {}).get("primaryCompletionDateStruct", {}).get("date", "")
            enrollment = ps.get("designModule", {}).get("enrollmentInfo", {}).get("count")
            results.append({
                "nctId": nct_id,
                "title": title,
                "phases": phases,
                "status": status,
                "startDate": start_date,
                "primaryCompletionDate": completion_date,
                "enrollment": enrollment
            })
        return results, None
    except Exception as e:
        return [], str(e)

# Read xlsx
wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws = wb["Metastatic + PD-L1 expression"]

drug_rows = []
for row in ws.iter_rows(min_row=3, values_only=True):
    name = str(row[0]).strip() if row[0] else ""
    if name and name.lower() != "none":
        drug_rows.append(name)

unique_drugs = sorted(set(drug_rows))

# Skip placeholders
PLACEHOLDERS = {"Alternate NTRK inhibitor", "Platinum-doublet chemotherapy regimens",
                "Gemcitabine monotherapy", "Pemetrexed monotherapy"}

mapping = {}
print(f"Fetching NCT IDs for {len(unique_drugs)} drugs...\n")

for i, drug in enumerate(unique_drugs, 1):
    if drug in PLACEHOLDERS:
        mapping[drug] = {"searchTerm": drug, "note": "Placeholder/generic — no specific trial", "trials": []}
        print(f"[{i}/{len(unique_drugs)}] SKIP {drug} (placeholder)")
        continue

    search_term = normalize_term(drug)
    print(f"[{i}/{len(unique_drugs)}] {drug} -> search: '{search_term}'")
    results, err = search_nct(search_term)
    if err:
        print(f"  ERROR: {err}")
        mapping[drug] = {"searchTerm": search_term, "note": f"API error: {err}", "trials": []}
    else:
        mapping[drug] = {"searchTerm": search_term, "note": "", "trials": results}
        print(f"  Found {len(results)} trials")
        for t in results[:3]:
            print(f"    {t['nctId']} | {t['status']:20s} | {str(t['phases']):20s} | {t['title'][:80]}")
    time.sleep(0.3)

with open(OUTPUT_PATH, "w") as f:
    json.dump(mapping, f, indent=2, default=str)

print(f"\nDone! Saved to {OUTPUT_PATH}")
print(f"Total drugs: {len(unique_drugs)}")
print(f"Mapped: {sum(1 for v in mapping.values() if v['trials'])}")
print(f"Failed/empty: {sum(1 for v in mapping.values() if not v['trials'])}")
